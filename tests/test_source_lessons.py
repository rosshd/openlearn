"""Source courses teach concepts, including the real first-Continue path."""

import io
import json
import os
from pathlib import Path
import socket
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from unittest import mock

import pytest
from docx import Document

from openlearn import cli, source_context, tutor_service
from openlearn.web.services import OpenLearnWebServices, _present_response
from tests.test_source_context import consent, web_request

pytest_plugins = ["tests.test_source_context"]


@pytest.fixture
def source_course(course):
    topic, record = course
    metadata, body = cli.parse_topic(topic.path.read_text())
    metadata["web_source_start"] = True
    cli.write_text_atomic(topic.path, cli.format_topic(metadata, body))
    return cli.read_topic(topic.slug), record


def plan(record):
    return json.dumps({"concepts": [
        {"label": "Stack order", "objective": "Explain last-in, first-out order.", "source_ids": [record.source_id]},
        {"label": "Stack operations", "objective": "Predict the result of pushes and pops.", "source_ids": [record.source_id]},
    ]})


def lesson(label="Stack order"):
    return (
        "**Lesson:**\nA stack removes its most recently added item first.\n\n"
        "For example, placing a blue plate on a red plate means the blue plate comes off first."
        f"\n<!-- covered: {label} -->\n<!-- focus: {label} -->"
    )


def submit(service, request, slug):
    preview = service.preview_source_turn(slug, request)
    assert preview["ok"], preview
    request.source_approval = preview["approval"]
    with ThreadPoolExecutor(max_workers=1) as executor, mock.patch.object(tutor_service, "_EXECUTOR", executor):
        service.submit_turn(slug, request)
    return tutor_service.operation_status(slug, request.submission_id)


@pytest.mark.parametrize("intent,text", [("next", ""), ("question", "Teach me these concepts")])
def test_first_source_turn_plans_and_teaches_atomically(source_course, intent, text):
    topic, record = source_course
    service = OpenLearnWebServices()
    request = web_request(topic, intent=intent, text=text)
    with mock.patch.object(cli, "call_openai", return_value=plan(record)) as planner, mock.patch.object(cli, "call_openai_streaming", return_value=lesson()) as tutor:
        operation = submit(service, request, topic.slug)
    assert operation.status == "committed", operation.error_message
    assert planner.call_count == tutor.call_count == 1
    saved = cli.read_topic(topic.slug)
    assert saved.metadata["course_started"]
    assert [unit["title"] for unit in saved.metadata["course_units"]] == ["Stack order", "Stack operations"]
    assert saved.metadata["slide_coverage"] == {"1:1": ["Stack order"]}
    assert "Source excerpts provided:" not in saved.body
    events = cli.load_event_log(cli.topic_events_path(topic.slug))
    assert any(event["event_type"] == "source_context_used" and record.source_id in str(event["data"]) for event in events)
    assert "PRIVATE_GOAL" not in str(planner.call_args)
    assert "PRIVATE_GOAL" not in str(tutor.call_args)


def test_next_source_concept_reuses_saved_plan_and_records_coverage(source_course):
    topic, record = source_course
    service = OpenLearnWebServices()
    with mock.patch.object(cli, "call_openai", return_value=plan(record)) as planner, mock.patch.object(cli, "call_openai_streaming", side_effect=[lesson(), lesson("Stack operations")]):
        assert submit(service, web_request(topic, intent="next", text=""), topic.slug).status == "committed"
        current = cli.read_topic(topic.slug)
        result = submit(service, web_request(current, intent="next", text=""), topic.slug)
    assert result.status == "committed", result.error_message
    assert planner.call_count == 1
    saved = cli.read_topic(topic.slug)
    assert saved.metadata["current_unit"] == 2
    assert saved.metadata["slide_coverage"] == {"1:1": ["Stack order"], "2:1": ["Stack operations"]}


def test_observed_assignment_answer_is_repaired_before_it_is_saved(source_course):
    topic, record = source_course
    bad = (
        "**Lesson:**\nThe classwork asks for Single, Average, and Complete linkage.\n\n"
        "The matrix is missing, so I can't compute the clusters.\n\n"
        "Action: Reply with the matrix or say skip question 1."
    )
    with mock.patch.object(cli, "call_openai", return_value=plan(record)), mock.patch.object(cli, "call_openai_streaming", side_effect=[bad, lesson()]) as tutor:
        result = submit(OpenLearnWebServices(), web_request(topic, intent="next", text=""), topic.slug)
    assert result.status == "committed", result.error_message
    assert tutor.call_count == 2
    assert "can't compute" not in cli.read_topic(topic.slug).body


def test_failed_lesson_does_not_save_partial_plan(source_course):
    topic, record = source_course
    with mock.patch.object(cli, "call_openai", return_value=plan(record)), mock.patch.object(cli, "call_openai_streaming", return_value="**Next:**\nGive me the matrix."):
        with pytest.raises(cli.OpenLearnError):
            cli.ask_topic(topic.slug, "Continue to the next useful concept.", source_mode=True, input_func=consent, output_func=lambda _: None)
    saved = cli.read_topic(topic.slug)
    assert not saved.metadata.get("source_learning_plan")
    assert not saved.metadata.get("course_units")


def test_planning_is_disclosed_before_approval_and_cancel_sends_nothing(source_course):
    topic, _ = source_course
    before = topic.path.read_bytes(), cli.load_state(topic.slug)
    with mock.patch.object(cli, "call_openai") as planner, mock.patch.object(cli, "call_openai_streaming") as tutor:
        preview = OpenLearnWebServices().preview_source_turn(topic.slug, web_request(topic, intent="next", text=""))
        assert preview["ok"]
        assert "concept" in preview["preview"].lower()
        assert "planning" in preview["disclosure"].lower()
        with pytest.raises(cli.OpenLearnError, match="cancelled"):
            cli.ask_topic(topic.slug, "Continue to the next useful concept.", source_mode=True, input_func=lambda _: "", output_func=lambda _: None)
    planner.assert_not_called()
    tutor.assert_not_called()
    assert (topic.path.read_bytes(), cli.load_state(topic.slug)) == before


@pytest.mark.parametrize("raw", ['{"concepts": []}', '{"concepts": [{"label": "Question 1", "objective": "Solve it", "source_ids": ["file:invented"]}]}'])
def test_invalid_source_plan_is_not_saved(source_course, raw):
    topic, _ = source_course
    with mock.patch.object(cli, "call_openai", return_value=raw), mock.patch.object(cli, "call_openai_streaming") as tutor:
        with pytest.raises(cli.OpenLearnError, match="concept sequence"):
            cli.ask_topic(topic.slug, "Continue to the next useful concept.", source_mode=True, input_func=consent, output_func=lambda _: None)
    tutor.assert_not_called()
    assert not cli.read_topic(topic.slug).metadata.get("source_learning_plan")


def test_real_emitted_ledger_is_hidden_without_mutating_history(course):
    topic, _ = course
    context = source_context.snapshot(topic, "stack", source_context.APPROVED_MODEL, opted_in=True)
    body = lesson() + context.ledger
    _, blocks = _present_response(body)
    assert "Source excerpts provided:" not in str(blocks)
    assert "checksum" not in str(blocks)
    assert "Bounded extracted text" not in str(blocks)
    assert "For example" in str(blocks)


def test_docx_tables_are_extracted_in_document_order():
    document = Document()
    document.add_paragraph("Distances:")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "0", "7"
    table.cell(1, 0).text, table.cell(1, 1).text = "7", "0"
    document.add_paragraph("Compare linkage methods.")
    stream = io.BytesIO()
    document.save(stream)
    text = cli._extract_docx_bytes(stream.getvalue(), "synthetic.docx")
    assert text == "Distances:\n0 | 7\n7 | 0\nCompare linkage methods."


def test_plan_freshness_is_checked_after_planning(source_course):
    topic, record = source_course
    def planner(*_args, **_kwargs):
        (cli.topic_context_dir(topic.slug) / record.context_file).write_text("Changed text")
        return plan(record)
    with mock.patch.object(cli, "call_openai", side_effect=planner), mock.patch.object(cli, "call_openai_streaming") as tutor:
        with pytest.raises(cli.OpenLearnError):
            cli.ask_topic(topic.slug, "Continue to the next useful concept.", source_mode=True, input_func=consent, output_func=lambda _: None)
    tutor.assert_not_called()
    assert not cli.read_topic(topic.slug).metadata.get("source_learning_plan")


def test_practice_and_side_question_preserve_plan_and_check(source_course):
    topic, record = source_course
    service = OpenLearnWebServices()
    check = "**Check:**\nAfter pushing a red plate then a blue plate, which plate comes off first?"
    with mock.patch.object(cli, "call_openai", return_value=plan(record)) as planner, mock.patch.object(cli, "call_openai_streaming", side_effect=[lesson(), check, "**Lesson:**\nLast-in, first-out means the newest item leaves first."]):
        assert submit(service, web_request(topic, intent="next", text=""), topic.slug).status == "committed"
        topic = cli.read_topic(topic.slug)
        result = submit(service, web_request(topic, intent="practice", text=""), topic.slug)
        assert result.status == "committed", result.error_message
        before = cli.read_topic(topic.slug).metadata
        assert submit(service, web_request(cli.read_topic(topic.slug), intent="question", text="Explain last-in, first-out"), topic.slug).status == "committed"
    after = cli.read_topic(topic.slug).metadata
    assert after["pending_question"] == before["pending_question"]
    assert after["source_learning_plan"] == before["source_learning_plan"]
    assert after["current_unit"] == before["current_unit"]
    assert planner.call_count == 1


@pytest.mark.parametrize("intent,text", [("next", ""), ("question", "Teach me these concepts")])
@pytest.mark.skipif(os.environ.get("OPENLEARN_BROWSER_TEST") != "1", reason="requires installed Playwright browser")
def test_source_lesson_browser_journey(source_course, intent, text):
    from playwright.sync_api import sync_playwright, expect
    import uvicorn
    from openlearn.web import create_app

    topic, record = source_course
    with socket.socket() as listener:
        listener.bind(("127.0.0.1", 0))
        port = listener.getsockname()[1]
    app = create_app(testing=True)
    server = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=port, log_level="error"))
    thread = threading.Thread(target=server.run, daemon=True)
    thread.start()
    deadline = time.monotonic() + 10
    while not server.started and time.monotonic() < deadline:
        time.sleep(0.02)
    assert server.started
    origin = f"http://127.0.0.1:{port}"
    base = origin + app.state.security.url_namespace
    check = "**Check:**\nAfter pushing red then blue onto the stack, which item is removed first?"
    try:
        with mock.patch.object(cli, "call_openai", return_value=plan(record)) as planner, mock.patch.object(cli, "call_openai_streaming", side_effect=[lesson(), lesson("Stack operations"), check]), sync_playwright() as runtime:
            browser = runtime.chromium.launch()
            page = browser.new_page(viewport={"width": 1280, "height": 900})
            page.goto(f"{origin}/?access_token={app.state.security.access_token}")
            page.goto(f"{base}/courses/{topic.slug}")
            expect(page.get_by_role("heading", name="Ready to learn", exact=True)).to_be_visible()
            before = topic.path.read_bytes(), cli.load_state(topic.slug)
            if intent == "question":
                page.locator('[data-tool-open="chat"]').first.click()
                page.locator("#chat-question").fill(text)
                page.locator("[data-chat-submit]").click()
            else:
                page.get_by_role("button", name="Continue", exact=True).click()
            expect(page.locator("[data-source-preview]")).to_be_visible()
            page.locator("[data-source-cancel]").click()
            assert (topic.path.read_bytes(), cli.load_state(topic.slug)) == before
            if intent == "question":
                page.locator("[data-chat-submit]").click()
            else:
                page.get_by_role("button", name="Continue", exact=True).click()
            page.locator("[data-source-send]").click()
            expect(page.locator("[data-current-move]")).to_contain_text("blue plate", timeout=10000)
            page.reload()
            expect(page.locator("[data-current-move]")).to_contain_text("blue plate")
            assert page.locator("[data-source-mode]").is_checked()
            for fragment in ("checksum", "Source excerpts provided:", record.source_id):
                assert fragment not in page.locator("[data-current-move]").inner_text()
            page.get_by_role("button", name="Continue", exact=True).click()
            page.locator("[data-source-send]").click()
            expect(page.locator("[data-current-move]")).to_contain_text("Stack operations", timeout=10000)
            page.get_by_role("button", name="Continue", exact=True).click()
            page.locator("[data-source-send]").click()
            expect(page.locator("[data-current-move]")).to_contain_text("which item", timeout=10000)
            page.reload()
            expect(page.locator("[data-current-move]")).to_contain_text("which item")
            assert planner.call_count == 1
            screenshots = Path(__file__).resolve().parents[1] / ".artifacts/source-concepts"
            screenshots.mkdir(parents=True, exist_ok=True)
            page.screenshot(path=str(screenshots / f"{intent}-check.png"))
            browser.close()
    finally:
        server.should_exit = True
        thread.join(timeout=10)


def test_first_practice_plans_without_forcing_a_lesson(source_course):
    topic, record = source_course
    check = "**Check:**\nAfter pushing red then blue onto a stack, which item leaves first?"
    with mock.patch.object(cli, "call_openai", return_value=plan(record)), mock.patch.object(cli, "call_openai_streaming", return_value=check):
        result = submit(OpenLearnWebServices(), web_request(topic, intent="practice", text=""), topic.slug)
    assert result.status == "committed", result.error_message
    saved = cli.read_topic(topic.slug)
    assert saved.metadata["source_learning_plan"]
    assert saved.metadata["pending_question"]
    assert not saved.metadata["slide_coverage"]


def test_source_audit_response_is_repaired_in_side_chat(source_course):
    topic, record = source_course
    good = "**Lesson:**\nThe newest stack item leaves first."
    audit = f"**Lesson:**\nsource ID {record.source_id}; original checksum abc; extracted text lines 1-2"
    with mock.patch.object(cli, "call_openai", return_value=plan(record)), mock.patch.object(cli, "call_openai_streaming", side_effect=[lesson(), audit, good]) as tutor:
        service = OpenLearnWebServices()
        assert submit(service, web_request(topic, intent="next", text=""), topic.slug).status == "committed"
        result = submit(service, web_request(cli.read_topic(topic.slug), intent="question", text="Explain stacks"), topic.slug)
    assert result.status == "committed", result.error_message
    assert tutor.call_count == 3
    assert "source ID" not in cli.read_topic(topic.slug).body
